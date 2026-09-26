"""
Planilha de revisão da família (Fase B da rodada 2): scripts/exportar_revisao.py
e scripts/importar_revisao.py, com o módulo comum revisao.py.

Teste de ida e volta numa CÓPIA dos textos: exportar → editar a planilha →
importar → os YAMLs têm exatamente as mudanças pedidas e o resto intacto
(inclusive os comentários).
"""
import importlib.util
import shutil

import pytest
import yaml
from openpyxl import load_workbook

from test_app import RAIZ

import montar_texto_ocidental as mt  # noqa: E402
import revisao  # noqa: E402


def _script(nome):
    spec = importlib.util.spec_from_file_location(nome, RAIZ / "scripts" / f"{nome}.py")
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


exportar, importar = _script("exportar_revisao"), _script("importar_revisao")


@pytest.fixture
def base_copiada(tmp_path, monkeypatch):
    """Os YAMLs numa pasta temporária: o teste nunca mexe nos de verdade."""
    pasta = tmp_path / "textos"
    shutil.copytree(mt.PASTA, pasta)
    monkeypatch.setattr(mt, "PASTA", pasta)
    mt.base.cache_clear()
    yield pasta
    mt.base.cache_clear()


def _linha(ws, id_):
    for r in range(2, ws.max_row + 1):
        if ws.cell(row=r, column=1).value == id_:
            return r
    raise KeyError(id_)


def _texto_ok(palavras=70):
    return ("Com Vênus em Libra, você valoriza a delicadeza nos encontros e percebe quando o clima "
            "pesa. " + " ".join(["Você costuma buscar acordos justos no dia a dia."] * 20)).split()[:palavras]


def test_exportar_tem_todas_as_abas_e_linhas(tmp_path, base_copiada):
    arq = exportar.gerar(tmp_path / "r.xlsx")
    wb = load_workbook(arq)
    assert wb.sheetnames[0] == "Como revisar"
    assert {"Mapa natal", "Sinastria"} <= set(wb.sheetnames)
    ws = wb["Mapa natal"]
    assert [c.value for c in ws[1]] == ["id", "onde aparece", "texto atual", "texto revisado",
                                         "aprovado?", "comentário"]
    total = sum(wb[a].max_row - 1 for a in wb.sheetnames[1:])
    assert total == len(revisao.itens())
    guia = " ".join(str(c.value or "") for c in wb["Como revisar"]["A"])
    for trecho in ("segunda pessoa", "60 a 120 palavras", "saúde", "morte", "dinheiro", "você vai"):
        assert trecho in guia


def test_ida_e_volta(tmp_path, base_copiada):
    antes = {a: (base_copiada / f"{a}.yaml").read_text(encoding="utf-8") for a in mt.ARQUIVOS}
    dados_antes = {a: yaml.safe_load(t) for a, t in antes.items()}
    arq = exportar.gerar(tmp_path / "r.xlsx")
    wb = load_workbook(arq)
    ws, ws_s = wb["Mapa natal"], wb["Sinastria"]
    novo_venus = " ".join(_texto_ok(70))
    # 1) texto novo + aprovado
    r = _linha(ws, "planetas_signos:venus/libra")
    ws.cell(row=r, column=4, value=novo_venus)
    ws.cell(row=r, column=5, value="sim")
    # 2) só aprovado (texto atual fica)
    ws.cell(row=_linha(ws, "ascendente:touro"), column=5, value="Sim")
    # 3) texto novo curto demais → recusado
    ws.cell(row=_linha(ws, "planetas_casas:marte/7"), column=4, value="Curto demais.")
    ws.cell(row=_linha(ws, "planetas_casas:marte/7"), column=5, value="sim")
    # 4) palavra proibida → recusado
    proibido = " ".join(_texto_ok(70)[:-3] + ["com", "sua", "saúde"])
    ws.cell(row=_linha(ws, "planetas_signos:sol/aries"), column=4, value=proibido)
    # 5) peça em formato compacto {revisado:..., texto:...} + aprovada
    ws.cell(row=_linha(ws, "pecas:temas/venus"), column=4, value="o seu jeito de gostar")
    ws.cell(row=_linha(ws, "pecas:temas/venus"), column=5, value="sim")
    # 6) sinastria: marcador {a} removido → recusado; comentário não muda nada
    r_s = _linha(ws_s, "sinastria:dimensoes/atracao/alta")
    ws_s.cell(row=r_s, column=4, value=ws_s.cell(row=r_s, column=3).value.replace("{a}", "Ana"))
    ws_s.cell(row=_linha(ws_s, "sinastria:casas/7"), column=6, value="gostei muito")
    wb.save(arq)

    alteracoes, recusados, desconhecidos = importar.planejar(importar.ler(arq))
    assert {a["id"] for a in alteracoes} == {"planetas_signos:venus/libra", "ascendente:touro",
                                             "pecas:temas/venus"}
    assert {r[0] for r in recusados} == {"planetas_casas:marte/7", "planetas_signos:sol/aries",
                                         "sinastria:dimensoes/atracao/alta"}
    assert desconhecidos == []

    # sem --gravar, nada muda
    assert all((base_copiada / f"{a}.yaml").read_text(encoding="utf-8") == antes[a] for a in mt.ARQUIVOS)

    revisao.gravar(alteracoes)
    depois = {a: yaml.safe_load((base_copiada / f"{a}.yaml").read_text(encoding="utf-8")) for a in mt.ARQUIVOS}
    assert " ".join(depois["planetas_signos"]["venus"]["libra"]["texto"].split()) == novo_venus
    assert depois["planetas_signos"]["venus"]["libra"]["revisado"] is True
    assert depois["ascendente"]["touro"]["revisado"] is True
    assert depois["ascendente"]["touro"]["texto"] == dados_antes["ascendente"]["touro"]["texto"]
    assert depois["pecas"]["temas"]["venus"] == {"revisado": True, "texto": "o seu jeito de gostar"}
    # recusados intactos
    assert depois["planetas_casas"]["marte"][7] == dados_antes["planetas_casas"]["marte"][7]
    assert depois["planetas_signos"]["sol"]["aries"] == dados_antes["planetas_signos"]["sol"]["aries"]
    # o resto, intacto: arquivos não tocados byte a byte, e comentários preservados
    for a in ("planetas_casas", "aspectos_pessoais", "sinastria"):
        assert (base_copiada / f"{a}.yaml").read_text(encoding="utf-8") == antes[a]
    for a in ("planetas_signos", "ascendente", "pecas"):
        comentarios = [l for l in antes[a].splitlines() if l.lstrip().startswith("#")]
        novo = (base_copiada / f"{a}.yaml").read_text(encoding="utf-8")
        assert all(c in novo for c in comentarios)
    # e o que o site lê depois é o texto novo
    assert mt.texto_signo("venus", "libra").split() == novo_venus.split()
    assert mt.contagem_de_revisao()["_total"][0] == 3


def test_id_desconhecido_e_ignorado(tmp_path, base_copiada):
    arq = exportar.gerar(tmp_path / "r.xlsx")
    wb = load_workbook(arq)
    wb["Mapa natal"].cell(row=2, column=1, value="planetas_signos:quiron/aries")
    wb.save(arq)
    _, _, desconhecidos = importar.planejar(importar.ler(arq))
    assert desconhecidos == ["planetas_signos:quiron/aries"]


def test_regras_de_tom_iguais_as_dos_testes():
    ok = " ".join(_texto_ok(70))
    assert revisao.problemas(ok, "planetas_signos", ("venus", "libra")) == []
    assert revisao.problemas("Você vai " + ok, "planetas_signos", ("venus", "libra"))
    assert revisao.problemas(" ".join(_texto_ok(130)), "planetas_signos", ("venus", "libra"))
    assert revisao.problemas("curta", "pecas", ("temas", "venus")) == []  # peça: sem limite


def test_gravar_no_tarot(base_copiada):
    """A frase das cartas é de uma linha só no YAML ({revisado, texto}); a gravação
    precisa funcionar nela e na leitura, sem mexer na carta vizinha."""
    antes = mt.base("tarot")["cartas"]["o_mago"]
    revisao.gravar([
        {"arquivo": "tarot", "caminho": ("cartas", "o_louco", "frase"), "texto": "Um começo novo.", "revisado": True},
        {"arquivo": "tarot", "caminho": ("cartas", "o_louco", "leitura"), "texto": None, "revisado": True},
        {"arquivo": "tarot", "caminho": ("conjunto", "maiores", "2"), "texto": None, "revisado": True},
    ])
    b = mt.base("tarot")
    assert b["cartas"]["o_louco"]["frase"] == {"revisado": True, "texto": "Um começo novo."}
    assert b["cartas"]["o_louco"]["leitura"]["revisado"] is True
    assert b["conjunto"]["maiores"][2]["revisado"] is True
    assert b["cartas"]["o_mago"] == antes
