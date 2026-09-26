"""
Gera a planilha de revisão dos textos da versão ocidental para a família do Pedro.

    python scripts/exportar_revisao.py                 # cria revisao-textos.xlsx
    python scripts/exportar_revisao.py outro-nome.xlsx

Uma aba "Como revisar" e uma aba por produto (Mapa natal, Sinastria, ...), com
os textos ainda não revisados primeiro. Depois de preenchida, a planilha volta
para o site com scripts/importar_revisao.py.

A planilha é gerada: não vai para o repositório (.gitignore).
"""
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Alignment, Font, PatternFill  # noqa: E402
from openpyxl.worksheet.datavalidation import DataValidation  # noqa: E402

import revisao  # noqa: E402

COLUNAS = [("id", 24), ("onde aparece", 30), ("texto atual", 70), ("texto revisado", 70),
           ("aprovado?", 12), ("comentário", 36)]
ORDEM_ABAS = ["Mapa natal", "Sinastria", "Numerologia", "Tarot"]

COMO_REVISAR = [
    ("Como revisar os textos da Padmini", True),
    ("", False),
    ("Cada aba é um produto do site (Mapa natal, Sinastria...). Cada linha é um texto que "
     "aparece para o cliente. Os que ainda não foram revisados vêm primeiro.", False),
    ("", False),
    ("O que fazer em cada linha", True),
    ("1. Leia a coluna \"texto atual\" e a coluna \"onde aparece\" (ela diz em que parte do "
     "relatório o texto entra, por exemplo \"Mapa natal · Vênus em Libra\").", False),
    ("2. Se o texto está bom como está: escreva sim na coluna \"aprovado?\".", False),
    ("3. Se quiser mudar: escreva o texto novo INTEIRO na coluna \"texto revisado\" e sim em "
     "\"aprovado?\". Se ainda quiser pensar mais, deixe \"aprovado?\" em branco — o texto novo "
     "entra, mas continua na lista para revisar.", False),
    ("4. Dúvida ou sugestão para o Nicolas? Use a coluna \"comentário\". Ela não muda nada no site.", False),
    ("5. Não mexa nas colunas \"id\" e \"texto atual\": é por elas que o site sabe qual texto é qual.", False),
    ("", False),
    ("O tom da Padmini", True),
    ("• Falar direto com a pessoa, em segunda pessoa: \"você\" (no casal, \"vocês\").", False),
    ("• Tendência, não sentença: \"você tende a\", \"costuma\", \"pode\". Nunca \"você vai\", "
     "\"com certeza\", \"o seu destino é\".", False),
    ("• Concreto: um exemplo de comportamento do dia a dia (\"numa discussão, você...\").", False),
    ("• Nada de previsão de saúde, doença, morte, acidente, dinheiro, herança ou fortuna.", False),
    ("• Evitar frases que serviriam para qualquer pessoa (\"você é especial e sensível\").", False),
    ("", False),
    ("Tamanho", True),
    ("• Textos do mapa, da sinastria e da numerologia: de 60 a 120 palavras.", False),
    ("• Cartas do tarot: de 40 a 90 palavras.", False),
    ("• Peças curtas (temas, frases de ligação, descrições): sem limite, mas curtas.", False),
    ("Um texto fora do tamanho ou com palavra proibida NÃO entra no site: ele volta para "
     "vocês com o motivo.", False),
    ("", False),
    ("Palavras entre chaves", True),
    ("Alguns textos têm {a}, {b}, {nome}, {quem}, {de_quem}: o site troca pelo nome das pessoas. "
     "Mantenha exatamente as mesmas, no lugar que fizer sentido.", False),
]


def gerar(destino: Path) -> Path:
    wb = Workbook()
    guia = wb.active
    guia.title = "Como revisar"
    guia.column_dimensions["A"].width = 110
    for n, (linha, titulo) in enumerate(COMO_REVISAR, start=1):
        c = guia.cell(row=n, column=1, value=linha)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        if titulo:
            c.font = Font(bold=True, size=13 if n == 1 else 11)

    por_produto: dict[str, list[dict]] = {}
    for item in revisao.itens():
        por_produto.setdefault(item["produto"], []).append(item)
    abas = sorted(por_produto, key=lambda p: ORDEM_ABAS.index(p) if p in ORDEM_ABAS else 99)
    for produto in abas:
        ws = wb.create_sheet(produto)
        for col, (nome, largura) in enumerate(COLUNAS, start=1):
            c = ws.cell(row=1, column=col, value=nome)
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="3F2740")
            ws.column_dimensions[c.column_letter].width = largura
        ws.freeze_panes = "C2"
        validacao = DataValidation(type="list", formula1='"sim,não"', allow_blank=True)
        ws.add_data_validation(validacao)
        linhas = sorted(por_produto[produto], key=lambda i: i["revisado"])  # não revisados primeiro
        for r, item in enumerate(linhas, start=2):
            valores = [item["id"], item["onde"] + ("" if not item["revisado"] else " (já revisado)"),
                       item["texto"], "", "", ""]
            for col, v in enumerate(valores, start=1):
                c = ws.cell(row=r, column=col, value=v)
                c.alignment = Alignment(wrap_text=True, vertical="top")
            validacao.add(f"E{r}")
            ws.row_dimensions[r].height = max(30, min(400, 15 * (len(item["texto"]) // 60 + 1)))
    wb.save(destino)
    return destino


def main():
    destino = Path(sys.argv[1] if len(sys.argv) > 1 else "revisao-textos.xlsx")
    gerar(destino)
    total = len(revisao.itens())
    faltam = sum(1 for i in revisao.itens() if not i["revisado"])
    print(f"{destino}: {total} textos, {faltam} ainda sem revisão")


if __name__ == "__main__":
    main()
