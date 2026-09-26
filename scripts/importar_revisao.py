"""
Traz para o site a planilha revisada pela família (gerada por exportar_revisao.py).

    python scripts/importar_revisao.py revisao-textos.xlsx           # só mostra o que mudaria
    python scripts/importar_revisao.py revisao-textos.xlsx --gravar  # grava nos YAMLs

Para cada linha:
  - "texto revisado" preenchido → vira o texto do site;
  - "aprovado?" = sim → o texto (o novo, ou o atual) fica marcado revisado: true;
  - texto revisado fora do tamanho, com palavra proibida ou com os marcadores
    {a}/{b}/{nome}... trocados → a linha é RECUSADA inteira (nada dela é gravado)
    e aparece na lista de recusados, com o motivo.
Depois de gravar: rodar os testes e abrir um PR, como qualquer mudança.
"""
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from openpyxl import load_workbook  # noqa: E402

import revisao  # noqa: E402

SIM = {"sim", "s", "x", "ok", "yes", "aprovado"}


def ler(planilha: Path) -> list[dict]:
    wb = load_workbook(planilha, read_only=True, data_only=True)
    linhas = []
    for ws in wb.worksheets:
        cabecalho = None
        for valores in ws.iter_rows(values_only=True):
            if cabecalho is None:
                if valores and valores[0] == "id":
                    cabecalho = [str(v or "").strip() for v in valores]
                    continue
                break  # aba sem tabela (ex.: "Como revisar")
            d = dict(zip(cabecalho, valores))
            if d.get("id"):
                linhas.append({"aba": ws.title, **d})
    return linhas


def planejar(linhas: list[dict]) -> tuple[list[dict], list[tuple], list[str]]:
    """(alterações, recusados [(id, motivos)], ids desconhecidos)."""
    por_id = {i["id"]: i for i in revisao.itens()}
    alteracoes, recusados, desconhecidos = [], [], []
    for l in linhas:
        item = por_id.get(str(l["id"]).strip())
        if not item:
            desconhecidos.append(str(l["id"]))
            continue
        novo = " ".join(str(l.get("texto revisado") or "").split())
        aprovado = str(l.get("aprovado?") or "").strip().lower() in SIM
        if novo == item["texto"]:
            novo = ""
        if novo:
            erros = revisao.problemas(novo, item["arquivo"], item["caminho"], item["texto"])
            if erros:
                recusados.append((item["id"], erros))
                continue
        if not novo and not (aprovado and not item["revisado"]):
            continue  # nada a fazer nesta linha
        alteracoes.append({"id": item["id"], "arquivo": item["arquivo"], "caminho": item["caminho"],
                           "texto": novo or None, "revisado": True if aprovado else None,
                           "onde": item["onde"]})
    return alteracoes, recusados, desconhecidos


def main():
    if len(sys.argv) < 2:
        sys.exit(__doc__)
    planilha = Path(sys.argv[1])
    alteracoes, recusados, desconhecidos = planejar(ler(planilha))
    novos = [a for a in alteracoes if a["texto"]]
    aprovados = [a for a in alteracoes if a["revisado"]]
    print(f"{planilha}: {len(novos)} textos novos, {len(aprovados)} aprovados, "
          f"{len(recusados)} recusados, {len(desconhecidos)} ids desconhecidos\n")
    for a in alteracoes:
        o_que = "texto novo" + (" + aprovado" if a["revisado"] else "") if a["texto"] else "aprovado"
        print(f"  {o_que:22s} {a['onde']}")
    if recusados:
        print("\nRECUSADOS (nada destas linhas foi gravado):")
        for id_, erros in recusados:
            print(f"  {id_}: {'; '.join(erros)}")
    if desconhecidos:
        print("\nIDs que não existem mais na base (linha ignorada):", ", ".join(desconhecidos))
    if "--gravar" not in sys.argv:
        print("\nNada foi gravado. Para gravar: acrescente --gravar.")
        return
    if alteracoes:
        gravados = revisao.gravar(alteracoes)
        print("\nGravado em:", ", ".join(f"conteudo/ocidental/textos/{g}.yaml" for g in gravados))


if __name__ == "__main__":
    main()
